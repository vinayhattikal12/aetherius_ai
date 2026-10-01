import io
import re
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from backend.app.services.generators.document_models import DocumentSpec, TableData


class ExcelGenerator:
    """
    Microsoft Excel (.xlsx) spreadsheet generator.
    Produces clean spreadsheets with auto-adjusted column widths,
    bold headers, zebra row striping, and numeric formatting.
    """

    @classmethod
    def generate_excel_bytes(cls, spec: DocumentSpec) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = spec.title[:30] if spec.title else "Data Sheet"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        body_font = Font(name="Calibri", size=10, color="0F172A")
        
        thin_border_side = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(
            left=thin_border_side,
            right=thin_border_side,
            top=thin_border_side,
            bottom=thin_border_side
        )

        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_center = Alignment(horizontal="center", vertical="center")

        current_row = 1

        # 1. Document Title Header
        title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        ws.cell(row=current_row, column=1, value=spec.title).font = title_font
        current_row += 1

        if spec.subtitle:
            sub_font = Font(name="Calibri", size=10, italic=True, color="64748B")
            ws.cell(row=current_row, column=1, value=spec.subtitle).font = sub_font
            current_row += 1

        current_row += 1  # Blank spacing row

        # 2. Render Tables from Sections
        tables_found = False
        for s in spec.sections:
            if s.table and s.table.headers:
                tables_found = True
                if s.heading and s.heading != "Overview":
                    sec_font = Font(name="Calibri", size=12, bold=True, color="334155")
                    ws.cell(row=current_row, column=1, value=s.heading).font = sec_font
                    current_row += 1

                # Header Row
                for col_idx, h_text in enumerate(s.table.headers, 1):
                    c = ws.cell(row=current_row, column=col_idx, value=h_text)
                    c.fill = header_fill
                    c.font = header_font
                    c.alignment = align_center
                    c.border = cell_border
                current_row += 1

                # Data Rows
                for r_idx, row_vals in enumerate(s.table.rows):
                    fill = zebra_fill if r_idx % 2 == 1 else white_fill
                    for col_idx, val in enumerate(row_vals, 1):
                        clean_v = str(val).replace("**", "").strip()
                        
                        # Detect numeric vs text
                        numeric_val = None
                        if re.match(r"^-?\d+(\.\d+)?$", clean_v):
                            try:
                                numeric_val = float(clean_v) if "." in clean_v else int(clean_v)
                            except ValueError:
                                pass

                        c = ws.cell(row=current_row, column=col_idx, value=numeric_val if numeric_val is not None else clean_v)
                        c.fill = fill
                        c.font = body_font
                        c.border = cell_border
                        c.alignment = align_right if numeric_val is not None else align_left
                    current_row += 1

                current_row += 2  # Spacing between tables

        # If no explicit markdown tables found, write sections as structured text sheet
        if not tables_found:
            for s in spec.sections:
                if s.heading:
                    ws.cell(row=current_row, column=1, value=s.heading).font = Font(name="Calibri", size=12, bold=True, color="1E3A8A")
                    current_row += 1
                for p in s.content_paragraphs:
                    ws.cell(row=current_row, column=1, value=p).font = body_font
                    current_row += 1
                for b in s.bullet_points:
                    ws.cell(row=current_row, column=1, value=f"• {b}").font = body_font
                    current_row += 1
                current_row += 1

        # 3. Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
